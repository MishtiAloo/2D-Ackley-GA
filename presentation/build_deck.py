from pathlib import Path
from math import log10, exp
import json, math
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

HERE=Path(__file__).resolve().parent; ROOT=HERE.parent; A=HERE/'assets'
S=json.loads((HERE/'data/summary.json').read_text()); RUN=json.loads((HERE/'data/run.json').read_text())
P=Presentation(); P.slide_width=Inches(16); P.slide_height=Inches(9)
BG='F7F4EC'; INK='18343B'; TEAL='0A7EA4'; BLUE='2F88FF'; ORANGE='C4671A'; GOLD='AF821E'; MUTED='596B70'; LINE='D8DFD9'; WHITE='FFFFFF'; PALE='E9F1ED'; PINK='B8318F'
DATE='12 September 2026'; FONT='Aptos'; HEAD='Aptos Display'
FIG=0; SLIDES=[]; ICONPARTS={}; NOTE=[]
def rgb(c):return RGBColor.from_string(c)
def rect(s,x,y,w,h,fill=WHITE,line=None,r=.0,name=None):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    if line: sh.line.color.rgb=rgb(line);sh.line.width=Pt(1)
    else:sh.line.fill.background()
    if r:sh.adjustments[0]=r
    if name:sh.name=name
    return sh
def text(s,t,x,y,w,h=.5,size=22,color=INK,bold=False,align=None,name=None,font=FONT):
    assert size>=14
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    for i,line in enumerate(t.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
        p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
        p.space_before=Pt(0);p.space_after=Pt(4)
        if align is not None:p.alignment=align
    if name:sh.name=name
    return sh
def line(s,x1,y1,x2,y2,color=LINE,width=1.3,arrow=False):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(width)
    if arrow:
        tail=OxmlElement('a:tailEnd');tail.set('type','triangle');tail.set('w','sm');tail.set('len','sm');sh._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln').append(tail)
    return sh
def circle(s,x,y,d,fill=TEAL,linecolor=None):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    if linecolor:sh.line.color.rgb=rgb(linecolor)
    else:sh.line.fill.background()
    return sh
def pill(s,label,x,y,w,color=TEAL):
    rect(s,x,y,w,.42,color,r=.2);text(s,label,x+.1,y+.065,w-.2,.3,14,WHITE,True,PP_ALIGN.CENTER)
def new(title,section,sub=None):
    s=P.slides.add_slide(P.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
    rect(s,.65,.42,.4,.06,TEAL);text(s,section.upper(),1.2,.30,12,.3,14,TEAL,True)
    text(s,title,.65,.85,14.8,.8,36,INK,True,font=HEAD)
    if sub:text(s,sub,.68,1.67,14.5,.52,18,MUTED)
    line(s,.65,8.46,15.35,8.46)
    text(s,'2D ACKLEY  /  GENETIC ALGORITHM',.65,8.62,7,.25,14,MUTED)
    text(s,DATE,10.3,8.62,3.6,.25,14,MUTED,align=PP_ALIGN.RIGHT)
    text(s,f'{len(P.slides):02d}',14.55,8.58,.8,.35,18,TEAL,True,PP_ALIGN.RIGHT)
    trans=OxmlElement('p:transition');trans.set('spd','med');trans.append(OxmlElement('p:fade'));s._element.append(trans)
    SLIDES.append({'title':title,'section':section});return s
def note(s,t):
    s.notes_slide.notes_text_frame.text=t
def caption(s,name,desc,x,y,w):
    global FIG;FIG+=1
    text(s,f'Fig. {FIG:02d} — {name}. {desc}',x,y,w,.55,14,MUTED)
def picture(s,filename,x,y,w,h,name,desc):
    path=A/filename
    im=Image.open(path);iw,ih=im.size;ratio=min(w/iw,h/ih);pw,ph=iw*ratio,ih*ratio
    px=x+(w-pw)/2;py=y+(h-ph)/2
    sh=s.shapes.add_picture(str(path),Inches(px),Inches(py),width=Inches(pw),height=Inches(ph))
    sh.name='Figure: '+name
    caption(s,name,desc,x,y+h+.10,w)
    return sh
def svg_icon(s,stem,x,y,size=.72):
    src=next((ROOT/'icons').rglob(stem+'.svg'))
    sh=s.shapes.add_picture(str(A/('icon_'+stem+'.png')),Inches(x),Inches(y),Inches(size),Inches(size))
    # Native SVG plus PNG compatibility fallback.
    if stem not in ICONPARTS:
        part=Part(PackURI('/ppt/media/native_'+stem+'.svg'),'image/svg+xml',P.part.package,src.read_bytes());ICONPARTS[stem]=part
    rid=s.part.relate_to(ICONPARTS[stem],RT.IMAGE)
    blip=sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
    extlst=OxmlElement('a:extLst');ext=OxmlElement('a:ext');ext.set('uri','{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
    svg=etree.Element('{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip',nsmap={'asvg':'http://schemas.microsoft.com/office/drawing/2016/SVG/main'})
    svg.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed',rid);ext.append(svg);extlst.append(ext);blip.append(extlst)
    sh.name='Icon: '+stem;return sh
def iconcard(s,stem,title,desc,x,y,w=4.65,h=2.1):
    rect(s,x,y,w,h,WHITE,r=.06)
    svg_icon(s,stem,x+.23,y+.22,.65)
    text(s,title,x+.23,y+1.0,w-.46,.4,24,INK,True)
    text(s,desc,x+.23,y+1.49,w-.46,h-1.5,17,MUTED)
def stat(s,value,label,x,y,w=3,color=TEAL,desc=None):
    text(s,value,x,y,w,.85,42,color,True,font=HEAD)
    text(s,label,x,y+.91,w,.6,20,INK,True)
    if desc:text(s,desc,x,y+1.53,w,.8,17,MUTED)
def callout(s,label,body,x,y,w=4.4,color=TEAL,idx=1):
    rect(s,x,y,.06,1.14,color)
    text(s,label,x+.2,y,w-.2,.45,23,color,True,name=f'anim_{idx}_title')
    text(s,body,x+.2,y+.52,w-.2,.8,20,INK,name=f'anim_{idx}_body')
def gene(s,vals,x,y,colors=(TEAL,PINK),label=None):
    if label:text(s,label,x-1.3,y+.14,1.2,.4,20,MUTED)
    for i,v in enumerate(vals):
        rect(s,x+i*1.85,y,1.7,.67,WHITE,colors[i],r=.07)
        text(s,str(v),x+i*1.85+.08,y+.13,1.54,.4,25,colors[i],True,PP_ALIGN.CENTER)
def stepbar(s,index):
    labels=['Init','Evaluate','Select','Cross','Mutate','Elite','Commit']
    for i,l in enumerate(labels):
        x=.65+i*2.11
        rect(s,x,2.15,1.98,.40,TEAL if i==index else PALE,r=.12)
        text(s,f'{i+1}  {l}',x+.04,2.205,1.9,.3,14,WHITE if i==index else MUTED,True,PP_ALIGN.CENTER)
def step_slide(title,index,filename,figdesc,items,notes):
    s=new(title,'Walkthrough · seed 42','One run throughout: 50 individuals · Pc 0.80 · Pm 0.05 · σ 1.0')
    stepbar(s,index)
    picture(s,filename,.65,2.85,9.15,4.65,title.split(' · ')[-1],figdesc)
    for j,(a,b) in enumerate(items):callout(s,a,b,10.3,3.0+j*1.6,4.9,idx=j+1)
    note(s,notes+'\nSource: ga_ackley.py; ga-web/engine.py; presentation/data/run.json. UI row identifiers are zero-based. Captures use the actual light-theme frontend with larger presentation text and selected rows where labelled.')
    return s

# 01 INTRO
s=new('Searching for the minimum','Project presentation')
text(s,'A genetic algorithm\nfor the 2D Ackley function',.7,2.25,6.7,1.9,39,INK,True,font=HEAD)
text(s,'From random points to a solution\nnear (0, 0).',.72,4.45,6.0,1.2,26,MUTED)
pill(s,'PYTHON + WEB VISUALIZER',.72,6.0,4.0)
text(s,'A visual, reproducible walkthrough',.72,6.65,6.2,.5,20,INK)
picture(s,'landscape_3d.png',7.1,2.18,8.25,5.65,'Ackley landscape','White-mode frontend; the central basin contains the known optimum.')
note(s,'Introduce the project as a population-based search for a known test-function minimum. The screenshot is the actual project frontend in light mode. The deck follows the default seed-42 run. No presenter or institution name was supplied, so none has been invented.')

# 02 CONTENTS
s=new('The route through the project','Contents','Problem → design choices → a real run → evidence')
contents=[('global-minimum-target','01  The problem','GA basics and the Ackley landscape'),('parameter-controls','02  The settings','Parameters and a tuning experiment'),('chromosome-dna','03  The methods','How each operator works—and why'),('numbered-stages','04  The walkthrough','Seven steps in the white frontend'),('convergence-curve','05  The evidence','Convergence, spikes, and final result'),('verdict-check','06  The takeaway','Conclusion and references')]
for i,(ic,t,d) in enumerate(contents):
    # Use the actual available target icon filename.
    if ic=='global-minimum-target':ic='global-minimum-target' if list((ROOT/'icons').rglob(ic+'.svg')) else 'global-minimum'
    iconcard(s,ic,t,d,.65+(i%3)*5.0,2.4+(i//3)*2.75,4.7,2.42)
note(s,'Agenda. Icons are from the supplied icons directory; each has its concept name and description beneath it.')

# 03 GA DEFINITION
s=new('GA: improve a population through generations','The idea','A genetic algorithm searches by selecting, combining, and randomly changing candidate solutions.')
iconcard(s,'01-population','Population','Several candidate answers at once.',.65,2.5,4.5,2.5)
iconcard(s,'03-selection-roulette','Selection','Better candidates get more chances.',5.75,2.5,4.5,2.5)
iconcard(s,'07-new-generation','Next generation','Keep useful changes; repeat.',10.85,2.5,4.5,2.5)
line(s,5.22,3.65,5.65,3.65,TEAL,2,True);line(s,10.32,3.65,10.75,3.65,TEAL,2,True)
rect(s,.65,5.65,14.7,1.66,PALE,r=.05)
text(s,'In this project',.95,5.92,3,.5,22,TEAL,True)
text(s,'One chromosome = [x₁, x₂]     •     One score = Ackley f(x₁, x₂)',4.0,5.95,10.8,.6,24,INK)
text(s,'Lower objective values are better.',4.0,6.63,10.8,.45,21,MUTED)
note(s,'A simple GA definition: evolve a collection of answers using biased reproduction and random variation. Real-valued chromosomes hold the two coordinates directly. The objective is minimized. General background: Holland, Adaptation in Natural and Artificial Systems, MIT Press, 1992 edition. https://doi.org/10.7551/mitpress/1090.001.0001')

# 04 PROBLEM
s=new('The problem: find the bottom of a rippled basin','Problem statement','Minimize Ackley over x₁, x₂ ∈ [−5, 5]. The known global minimum is f(0, 0) = 0.')
picture(s,'landscape_3d.png',.65,2.3,8.75,5.38,'Search landscape','White-mode 3D view; many local basins surround the origin.')
text(s,'min f(x₁, x₂)',9.9,2.55,5.1,.8,35,TEAL,True)
text(s,'−20 exp[−0.2 √((x₁²+x₂²)/2)]\n− exp[(cos 2πx₁ + cos 2πx₂)/2]\n+ 20 + e',9.9,3.47,5.2,1.6,21,INK,font='Cambria Math')
callout(s,'Why this is challenging','Local valleys can trap the search before it reaches the central basin.',9.9,5.55,5.2,idx=1)
note(s,'The implementation uses a=20, b=0.2, c=2π with a restricted domain [-5,5]^2. The commonly cited benchmark domain is larger; the project deliberately uses this smaller search square. The optimum remains (0,0). Source: ga_ackley.py, CONFIG and ackley; Surjanovic and Bingham, SFU, https://www.sfu.ca/~ssurjano/ackley.html .')

# 05 SETTINGS
s=new('The baseline settings balance search and preservation','GA settings','These are the project defaults used in every walkthrough and result slide.')
settings=[('50','individuals','Variety with a small evaluation budget.'),('0.80','crossover / pair','Recombine most pairs; keep some copies.'),('0.05','mutation / gene','About 5 changes among 100 genes.'),('1.0','Gaussian σ','Steps span 10% of the coordinate range.'),('1','elite','Protect progress while most slots evolve.'),('42','random seed','Replay the same random decisions.')]
for i,(v,l,d) in enumerate(settings):
    x=.65+(i%3)*5.0;y=2.4+(i//3)*2.75
    rect(s,x,y,4.7,2.5,WHITE,r=.05);stat(s,v,l,x+.25,y+.2,4.15,desc=d)
note(s,'Defaults: n_genes=2; bounds=-5,+5; N=50; Pc=.8 per pair; Pm=.05 per gene; elite_count=1; gaussian sigma=1; seed=42. Operators are also part of the assignment specification. These are practical baseline choices, not experimentally established global best settings. Expected mutated genes per generation = 50×2×.05 =5; actual count varies. Expected crossover pairs=25×.8=20. Stopping parameters appear on slide 11.')

# 06 TUNING
s=new('Tuning: smaller steps help, but can miss the basin','Parameter tuning','20 runs per setting · same seeds 42–61 · change one parameter · default stopping rules')
rows=[('Baseline',0),('Pc = 0.70',1),('Pm = 0.10',4),('σ = 0.10',5),('σ = 0.50',6),('N = 100',8)]
text(s,'SETTING',.8,2.4,3,.4,15,MUTED,True);text(s,'MEDIAN FINAL f(x)  ↓',4.0,2.4,5,.4,15,MUTED,True);text(s,'WORST RUN  ↓',10.1,2.4,3,.4,15,MUTED,True)
for i,(label,idx) in enumerate(rows):
    e=S['experiments'][idx];y=3.05+i*.63
    text(s,label,.8,y,3.05,.45,22,INK,idx==0)
    rect(s,4.0,y+.09,e['median']/.07*4.6,.28,TEAL if idx!=5 else GOLD,r=.05)
    text(s,f"{e['median']:.4f}",8.8,y,1.3,.4,20,TEAL,True)
    text(s,f"{e['max']:.4f}",10.2,y,2.0,.4,21,ORANGE if idx==5 else INK)
line(s,.8,6.93,15.1,6.93)
text(s,'σ = 0.10: strong median, weak worst case.',.8,7.16,7.2,.45,23,ORANGE,True)
text(s,'Tune exploration and precision together.',8.3,7.17,7.0,.45,23,TEAL,True)
text(s,'Exploratory comparison; larger N uses more evaluations. No setting reached f < 0.001 in these 20 runs.',.8,7.84,14.6,.37,14,MUTED)
note(s,'New experiment generated by presentation/build_data.py with the existing run_ga function. Each setting uses seeds 42 through 61, max_generations=100 and stagnation_limit=30; all other settings remain baseline. Nine configurations (180 runs) are stored in summary.json; six representative rows shown. Median and worst final best objective are reported because sigma=.1 has a good median but a heavy poor-run tail (mean .64875, max2.58022). This does not establish statistical superiority, and the 100-individual comparison is not an equal-evaluation-budget study. Baseline remains unchanged for the walkthrough. Data: presentation/data/summary.json.')

# 07 INITIALIZATION METHOD
s=new('Initialization: store real numbers and spread the search','Methods · encoding & initialization','Each coordinate is sampled independently and uniformly from [−5, 5].')
svg_icon(s,'random-initialization',.95,2.7,.8)
text(s,'Uniform initialization',.95,3.68,5.4,.5,26,INK,True)
text(s,'Equal sampling density across the domain.',.95,4.29,5.1,1,23,MUTED)
text(s,'Why: broad starting coverage,\nwith no preferred region.',.95,5.7,5.5,1.0,24,TEAL,True)
rect(s,7.0,2.65,8.0,4.75,WHITE,r=.045)
text(s,'One chromosome',7.4,2.95,6,.5,25,INK,True)
gene(s,['x₁','x₂'],8.5,3.87,colors=(TEAL,TEAL))
text(s,'Each gene ∼ U(−5, 5)',8.05,4.86,5.8,.6,27,INK,font='Cambria Math')
line(s,8.0,6.3,14.0,6.3,TEAL,2)
for xx,l in [(8,'−5'),(11,'0'),(14,'5')]:line(s,xx,6.16,xx,6.46,TEAL,2);text(s,l,xx-.4,6.62,.8,.4,19,MUTED,align=PP_ALIGN.CENTER)
for xx in [8.7,9.6,10.3,11.4,12.6,13.2]:circle(s,xx,6.21,.18,TEAL)
note(s,'Source: create_individual and create_population in ga_ackley.py. Real-valued encoding is direct: no binary decoding and no discretization step. Uniform independent draws are unbiased over the bounded search square, though finite samples do not guarantee every area is covered. Initialization and real encoding follow the project specification.')

# 08 FITNESS & SELECTION
s=new('Selection: turn low cost into a higher chance','Methods · fitness & roulette','Fitness is recomputed relative to the worst individual in each generation.')
text(s,'Fᵢ = fworst − fᵢ',.85,2.6,6.2,.8,35,TEAL,True,font='Cambria Math')
text(s,'pᵢ = Fᵢ / ΣF',.85,3.57,6.2,.8,32,INK,font='Cambria Math')
text(s,'Toy scores',8.15,2.6,6.2,.5,23,INK,True)
toy=[('A',2,6,.6),('B',4,4,.4),('C',8,0,0)]
for i,(lab,cost,fit,prob) in enumerate(toy):
    y=3.37+i*.76;text(s,f'{lab}: f = {cost}',8.15,y,2.0,.4,23,INK)
    rect(s,10.4,y+.09,prob*4,.3,TEAL if i==0 else BLUE)
    text(s,f'{int(prob*100)}%',13.15,y,1.5,.4,23,TEAL,True)
rect(s,.65,5.63,14.7,1.95,WHITE,r=.05)
svg_icon(s,'selection-probability',.95,5.93,.65)
text(s,'Roulette sampling',.95,6.71,3.5,.45,21,INK,True)
text(s,'50 draws with replacement',5.0,5.96,9.8,.5,26,TEAL,True)
text(s,'Why: reward better candidates while keeping selection random.\nIf every fitness is zero, choose uniformly at random.',5.0,6.59,9.6,.85,21,INK)
note(s,'Toy example is explicitly separate from the real-run walkthrough: f=[2,4,8], worst=8, fitness=[6,4,0], probabilities=[.6,.4,0]. The baseline uses fitness-proportionate roulette with replacement to fill a pool of N=50. The highest objective has zero fitness; if all objectives tie and total fitness is zero, the implementation falls back to uniform random indices. Tournament selection is an available alternative but is not used in the baseline. Source: objective_to_fitness, roulette_wheel_selection in ga_ackley.py.')

# 09 CROSSOVER METHOD
s=new('Crossover: swap the second coordinate','Methods · one-point crossover','With two genes, there is exactly one cut point.')
gene(s,['A₁','A₂'],2.35,2.85,(TEAL,TEAL),'Parent A')
gene(s,['B₁','B₂'],2.35,4.0,(PINK,PINK),'Parent B')
line(s,6.6,3.67,8.8,3.67,TEAL,3,True)
text(s,'Pc = 0.80',6.5,2.92,2.3,.5,24,TEAL,True,PP_ALIGN.CENTER)
gene(s,['A₁','B₂'],10.25,2.85,(TEAL,PINK),'Child 1')
gene(s,['B₁','A₂'],10.25,4.0,(PINK,TEAL),'Child 2')
line(s,4.13,2.65,4.13,4.85,GOLD,2)
text(s,'cut',3.72,5.03,.85,.4,20,GOLD,True,PP_ALIGN.CENTER)
iconcard(s,'gene-swap','Why this method?','Simple coordinate recombination; preserves useful gene values.',.65,5.88,7.1,2.06)
rect(s,8.25,5.88,7.1,2.06,PALE,r=.06)
text(s,'What it cannot do',8.55,6.16,6.4,.5,25,INK,True)
text(s,'Create a new coordinate value.\nThat job belongs to mutation.',8.55,6.78,6.35,.9,23,MUTED)
note(s,'The assignment specifies 1-point crossover. With two coordinates, child1=[A.x1,B.x2] and child2=[B.x1,A.x2]. One random number is drawn per pair; if it exceeds Pc the parents are copied. Crossover therefore recombines existing coordinate values but does not synthesize intermediate real values; arithmetic crossover is an available alternative in the code. Source: one_point_crossover and do_crossover, ga_ackley.py.')

# 10 MUTATION METHOD
s=new('Mutation: introduce new coordinate values','Methods · Gaussian mutation & bounds','For each gene: mutate with Pm = 0.05; otherwise keep the original value.')
text(s,'x′ = x + ε',.85,2.7,6,.8,39,TEAL,True,font='Cambria Math')
text(s,'ε ∼ Normal(0, σ²),  σ = 1',.85,3.75,6.6,.7,29,INK,font='Cambria Math')
text(s,'Why Gaussian?',.85,5.16,6.5,.5,27,INK,True)
text(s,'Explore around the current point;\noccasionally take a larger step.',.85,5.83,6.35,1.1,24,MUTED)
rect(s,8.0,2.55,7.35,4.9,WHITE,r=.04)
# Editable Gaussian curve, no bitmap chart.
line(s,8.6,5.5,14.75,5.5,LINE,1.5)
pts=[]
for i in range(81):
    z=-3+6*i/80;pts.append((8.7+6*i/80,5.5-2.15*exp(-z*z/2)))
for a,b in zip(pts,pts[1:]):line(s,*a,*b,TEAL,2.7)
text(s,'−3σ',8.35,5.73,1.2,.4,18,MUTED);text(s,'0',11.45,5.73,.5,.4,18,MUTED);text(s,'+3σ',14.05,5.73,1.2,.4,18,MUTED)
text(s,'Out of bounds? Resample in [−5, 5].',8.55,6.46,6.4,.7,21,ORANGE,True)
note(s,'The Gaussian parameter is standard deviation sigma. The mathematical notation Normal(0,sigma²) uses variance as its second parameter; random.gauss(0,sigma) uses standard deviation. Every gene gets an independent Bernoulli mutation decision. Bounds handling resamples a fresh coordinate uniformly if the proposed value leaves [-5,5]. This avoids invalid chromosomes and boundary pile-up caused by clipping, but an out-of-bounds mutation becomes a global jump. Source: mutate_population, mutate_gene, repair_gene in ga_ackley.py.')

# 11 ELITISM STOPPING
s=new('Protect the best—and know when to stop','Methods · elitism & stopping','Two stopping checks limit wasted work; neither proves the exact optimum has been found.')
iconcard(s,'preserve-best','Keep one elite','Replace the worst child only if the saved elite is better.',.65,2.5,4.7,3.03)
iconcard(s,'max-generations','100 generations','A fixed upper bound on the run length.',5.65,2.5,4.7,3.03)
iconcard(s,'stagnation-pause','30 stagnant generations','Stop if no improvement > 10⁻¹² persists for 30 generations.',10.65,2.5,4.7,3.03)
for x,t in [(.95,'Why: preserve progress.'),(5.95,'Why: control computation.'),(10.95,'Why: end an unproductive run.')]:text(s,t,x,5.94,4.12,.88,22,TEAL,True)
rect(s,.65,7.12,14.7,.8,PALE,r=.05)
text(s,'f(x) < 0.001 is a reporting threshold in this project, not an early-stop rule.',.95,7.31,14.0,.4,21,INK)
note(s,'Source: get_elites, apply_elitism and run_ga in ga_ackley.py; GAEngine evaluate and commit stages. Elitism protects the best objective from worsening, though an elite need not be inserted if all children are at least as good. Max_generations=100 is the hard run budget. Stagnation resets only when best-ever improvement is greater than 1e-12; smaller improvements are still recorded but count as stagnant. The success verdict f<1e-3 is checked in the reporting function, not used as a stopping condition. The returned history evaluates each generation at loop entry; a final post-commit population is not an extra history point.')

# 12 ALGORITHM FLOW
s=new('The complete GA flow','Algorithm map','Initialization happens once; the remaining stages repeat.')
nodes=[('1','Initialize',1.05,2.7),('2','Evaluate',4.65,2.7),('3','Select',8.25,2.7),('4','Crossover',11.85,2.7),('7','Commit',4.65,5.55),('6','Elitism',8.25,5.55),('5','Mutate',11.85,5.55)]
for n,l,x,y in nodes:
    rect(s,x,y,2.7,1.15,WHITE,LINE,r=.08);circle(s,x+.19,y+.22,.64,TEAL);text(s,n,x+.2,y+.3,.62,.4,23,WHITE,True,PP_ALIGN.CENTER);text(s,l,x+.98,y+.38,1.62,.5,22,INK,True)
for a,b in [(3.78,4.55),(7.38,8.15),(10.98,11.75)]:line(s,a,3.27,b,3.27,TEAL,2,True)
line(s,13.2,3.9,13.2,5.45,TEAL,2,True)
line(s,11.75,6.12,11.05,6.12,TEAL,2,True);line(s,8.15,6.12,7.45,6.12,TEAL,2,True)
line(s,6.0,5.48,6.0,4.18,TEAL,2,True)
text(s,'repeat',6.3,4.69,1.6,.4,18,TEAL)
rect(s,.65,5.52,3.1,1.72,PALE,r=.06)
text(s,'STOP → report best',.9,5.82,2.6,.8,24,TEAL,True)
text(s,'Budget or stagnation',.9,6.7,2.65,.35,16,MUTED)
line(s,4.65,3.9,2.2,5.4,ORANGE,1.7,True)
text(s,'check criteria',2.05,4.13,2.65,.5,18,ORANGE)
text(s,'Evaluate → transform fitness → save elite → select → cross → mutate + repair → elitism → commit',.85,7.69,14.5,.5,18,MUTED)
note(s,'Conceptual flow matches the implementation sequence. At evaluation, compute objectives, histories, fitness and elites, then test stagnation. The web engine checks generation budget at commit. When continuing, commit makes offspring the next population and evaluation begins again. Initialization occurs before generation1. The next seven slides expand these steps with authentic frontend panels; the frontend itself numbers six stages after initialization.')

# 13–19 WALKTHROUGH
step_slide('1 · Initialize 50 candidate points',0,'initial_map.png','Generation 1, seed 42; points cover the search square.',[
 ('One point = one chromosome','The coordinates are the two genes.'),('Uniform random start','No candidate is placed at the origin by design.'),('Repeatable example','Seed 42 fixes the random stream.')],
 'The initial population contains 50 two-gene chromosomes drawn uniformly. The captured map is after the initial evaluation so the frontend can highlight the best and worst; point positions are exactly the initialized population. No breeding has yet occurred.')

step_slide('2 · Evaluate and rank the population',1,'evaluate.png','White-mode inspector; five highest-ranked rows shown.',[
 ('Best is individual #17','f = 5.4475 at (−1.2147, 0.5204).'),('Convert cost to fitness','13.4771 − 5.4475 = 8.0296.'),('Save the elite','Keep a protected copy before breeding.')],
 'Generation1 has best objective5.447494814976, mean9.979817006302 and worst13.477089228593. Individual #17 has fitness8.029594413617. Total fitness174.863611114545 gives #17 a draw probability of approximately4.592%. The bars in this inspector are scaled to the largest displayed fitness for visualization; use normalized total fitness for true roulette probability.')

step_slide('3 · Select the mating pool',2,'select.png','White-mode inspector; roulette tape and complete mating pool.',[
 ('Draw 50 parent slots','The first four IDs: #1, #35, #32, #25.'),('Duplicates are allowed','#37 is selected four times.'),('Better is more likely','A strong candidate is favoured, not guaranteed.')],
 'Actual generation1 roulette draws: spin2.0076129562 picks#1; spin126.0280199424 picks#35; spin119.2063368622 picks#32; spin93.8965710367 picks#25. Sampling uses replacement. Most picked#37 appears4 times. The fitness tape and full pool are genuine frontend content. Some rows of the frequency summary are hidden solely to make a readable excerpt.')

step_slide('4 · Crossover creates two children',3,'crossover.png','White-mode inspector; pair 2 shown from 25 pairs.',[
 ('The random roll is 0.6499','0.6499 ≤ 0.80 → crossover happens.'),('Cut between x₁ and x₂','Parents #32 and #25 swap their tails.'),('19 of 25 pairs cross','The expected count is 20; actual counts vary.')],
 'Pair2 is selected for a simple illustration of one-point crossover. ParentA#32=[3.428519201898096,2.759999115462448]; ParentB#25=[-1.2981903288311738,-2.904929692285123]. Roll=.6498780576394535. Child2=[3.428519201898096,-2.904929692285123], child3=[-1.2981903288311738,2.759999115462448]. These exact values are from generation1, not invented toy numbers.')

step_slide('5 · Mutate selected genes',4,'mutate.png','White-mode inspector; three changed children shown.',[
 ('5 of 100 genes mutate','Each coordinate has its own 5% chance.'),('Child 5, first coordinate','0.5695 + 1.3945 = 1.9640.'),('Repair only if needed','All five mutations stay within [−5, 5].')],
 'Generation1 mutates child5 gene0, child8 gene1, child23 gene1, child29 gene1 and child48 gene0. Child5 roll=.0147578565, Gaussian delta1.39448230997, old.569497437746 -> new1.963979747715. None of the five proposals requires bounds repair. The UI excerpt shows children5,8,23; the three shown rows are not the complete mutation list.')

step_slide('6 · Insert the protected elite',5,'elitism.png','White-mode inspector; the actual generation-1 replacement.',[
 ('Saved elite: f = 5.4475','The best parent was protected before breeding.'),('Worst child: f = 13.1205','Child #47 is replaced by the better elite.'),('Population stays at 50','One replacement; no extra individual.')],
 'Actual event: elite[-1.214656227916465,.5204063127322698],f5.447494814976 replaces child47=[3.763676264726689,4.36654587712494],f13.120508066058. Other children may already be better than the elite; elitism is preservation, not a requirement to make the previous elite the next best.')

step_slide('7 · Commit and repeat',6,'commit.png','White-mode generation summary after the first breeding cycle.',[
 ('Children become generation 2','The best candidate now has f = 2.9802.'),('A better combination','New best: (0.8459, 0.0953).'),('Next: evaluate again','Update history, check stopping, and repeat.')],
 'Commit after generation1 changes population to the offspring. Best goes from5.447494814976 to2.980182401782 at [.8458599022354054,.09526293676464537]. The frontend reports the new best at commit, but its convergence history appends that value when generation2 is evaluated. This accounts for generation labels across screenshot panels.')

# 20 ANIMATION
s=new('Watch the population gather around the origin','Algorithm in motion','Animated white-mode frontend replay · selected evaluated generations from the same run')
picture(s,'evolution.gif',.65,2.23,9.1,5.49,'Population replay','10 snapshots, generations 1–100; loops in PowerPoint Slide Show.')
callout(s,'Early: broad movement','Selection and recombination rapidly reduce the objective.',10.25,2.85,4.8,idx=1)
callout(s,'Later: a compact cluster','New mutations still send some points away from the centre.',10.25,4.62,4.8,idx=2)
callout(s,'A persistent best point','Elitism keeps the strongest candidate safe.',10.25,6.39,4.8,idx=3)
note(s,'The embedded GIF uses actual frontend screenshots at evaluated generations1,2,5,10,17,30,47,67,77,100. It is a time-lapse, not every generation and not real elapsed time. The image animates in desktop PowerPoint Slide Show and loops; PDF exports show a static frame. The original GIF is provided in presentation/assets/evolution.gif.')

# 21 CONVERGENCE
s=new('Convergence: fast gains, long plateaus, a late drop','Results · convergence','Blue = best objective · orange = population average · logarithmic vertical scale')
picture(s,'convergence_log.png',.65,2.25,14.7,4.22,'Convergence curve','Actual white-mode frontend; 100 evaluated generations, seed 42.')
for x,v,l in [(1.0,'G1 → G5','5.4475 → 2.7318'),(6.05,'G47 → G76','Best stays at 0.10067'),(11.1,'G77 → G100','Best reaches 0.011695')]:
    text(s,v,x,7.12,4.4,.45,24,TEAL,True);text(s,l,x,7.65,4.4,.45,21,INK)
note(s,'History values are from the evaluated population at the start of each generation. Roughly49.96% of the total absolute best-value improvement occurs from generation1 to generation5. Later gains matter: the objective moves below1 at generation9 and reaches.0116950928353 at generation77. From generation47 through76 best remains .100666379838; the improvement at generation77 arrives just before 30 stagnant generations would trigger a stop. With one elite the best curve is nonincreasing. The orange mean oscillates because the population contains new, sometimes poor offspring. Log scale makes later improvements visible; it does not cause the spikes.')

# 22 SPIKE CAUSAL
s=new('Why did the average spike at generation 67?','Results · explain the curve','The generation-66 mutation stage produced the population evaluated at generation 67.')
vals=[('G66 population',.23126789337595507,TEAL),('After crossover',.17400936886974133,TEAL),('After mutation',.7876932534951682,ORANGE),('After elitism',.6761702778382442,GOLD)]
for i,(l,v,c) in enumerate(vals):
    x=.95+i*3.73
    rect(s,x,5.62-v*3.3,2.55,v*3.3,c,r=.015)
    text(s,f'{v:.4f}',x,5.72,2.65,.5,28,c,True,PP_ALIGN.CENTER)
    text(s,l,x-.1,6.35,2.85,.8,20,INK,True,PP_ALIGN.CENTER)
    if i<3:line(s,x+2.74,4.4,x+3.45,4.4,LINE,2,True)
line(s,.85,5.63,15.0,5.63,LINE,1)
text(s,'Mean f(x)',.85,2.35,3,.4,18,MUTED)
pill(s,'8 GENES MUTATED',7.8,2.36,3.0,ORANGE)
rect(s,.65,7.32,14.7,.67,PALE,r=.08)
text(s,'Best stays at 0.10067. Mutation worsened the mean; elitism removed the worst child.',.93,7.46,14.1,.43,21,INK)
note(s,'Exact stage decomposition: mean of generation66=.23126789337595507. After selection and crossover in generation66 the children have mean.17400936886974133. Eight genes mutate, taking the mean to.7876932534951682. Elitism replaces worst child31,f5.67681516268388 by elite f.10066637983769455, reducing mean to.6761702778382442. That is the generation67 history average. Net jump=.44490238446228914, the largest one-generation upward average jump in this run. Best remains .10066637983769455 at all three offspring stages. This is a directly measured explanation, not a generic assumption that all spikes are caused by mutation. Data: run.json, generations[65].')

# 23 FINAL RESULT
s=new('Final result: close to the known optimum','Results · seed 42','The run ends at the 100-generation limit, with 23 consecutive stagnant generations.')
rect(s,.65,2.5,7.1,4.6,WHITE,r=.045)
text(s,'Best point',1.0,2.85,6.4,.5,24,MUTED)
text(s,'(0.003973, 0.000317)',1.0,3.55,6.4,.8,34,INK,True)
text(s,'f(x) = 0.011695',1.0,4.65,6.4,.9,42,TEAL,True)
text(s,'Distance to origin: 0.003985',1.0,5.92,6.4,.6,25,INK)
stat(s,'99.79%','reduction from the initial best',8.55,2.85,6.4,desc='5.447495 → 0.011695')
rect(s,8.25,5.9,7.1,1.46,PALE,r=.06)
text(s,'Near the centre; still above 0.001.',8.55,6.18,6.4,.85,27,TEAL,True)
text(s,'The known answer is (0, 0), with f = 0.',.85,7.68,14.4,.45,22,MUTED)
note(s,'Exact best chromosome[.003972697250467887,.00031705775953021903]; best objective .011695092835325216; Euclidean distance .003985329229405449.100 evaluated generations; stop reason generation limit reached; stagnation23. Percent reduction from initial best =100×(5.447494814975997-.011695092835325216)/5.447494814975997 =99.7853%. The objective is above the project reporting threshold1e-3. Therefore this is a near-origin solution, not exact convergence to zero. Command-line and web histories match.')

# 24 WHY RESULT
s=new('Why did the run stop short of zero?','Results · interpretation','The final point is good, but the operators make very fine improvement increasingly unlikely.')
iconcard(s,'gene-swap','Limited recombination','One-point crossover can only reuse existing coordinates.',.65,2.5,4.7,3.15)
iconcard(s,'mutation-sigma','A fixed step size','σ = 1 is large compared with the final ≈0.004 distance.',5.65,2.5,4.7,3.15)
iconcard(s,'max-generations','A finite search budget','No better candidate appears after generation 77.',10.65,2.5,4.7,3.15)
rect(s,.65,6.25,14.7,1.56,PALE,r=.055)
text(s,'A useful next experiment',.95,6.53,5.0,.5,25,TEAL,True)
text(s,'Reduce σ later in the run, then test across many seeds.',6.0,6.57,8.95,.8,24,INK)
note(s,'Interpretation grounded in code and trace, not proof of a unique cause: one-point crossover creates no new coordinate values; mutation sigma remains1, about251 times the final Euclidean distance .003985; a narrow improvement near zero therefore requires a fortunate small perturbation. The best improves at generation77, and the remaining23 evaluated generations show no further progress before the100 cap. An adaptive mutation sigma is proposed future work; it is not implemented in this project. Tuning showed that a small constant sigma can have poor outliers, motivating a schedule rather than claiming that a universally small sigma is best.')

# 25 CONCLUSION
s=new('The GA found the right basin—and exposed its limits','Conclusion','A visual run makes both the progress and the remaining gap explainable.')
for i,(num,title,desc) in enumerate([('01','Selection directs the search','Fitter candidates receive more opportunities.'),('02','Variation explores; elitism protects','Average quality can fall while the best is preserved.'),('03','Tuning needs repeated evidence','Check typical quality, poor runs, and evaluation cost.')]):
    y=2.5+i*1.57;circle(s,.85,y,.8,TEAL);text(s,num,.86,y+.2,.77,.4,24,WHITE,True,PP_ALIGN.CENTER)
    text(s,title,2.02,y,12.4,.5,29,INK,True,name=f'anim_{i+1}_title');text(s,desc,2.02,y+.65,12.4,.6,23,MUTED,name=f'anim_{i+1}_body')
pill(s,'FINAL BEST  0.011695',.85,7.56,4.2)
text(s,'Near the optimum; further precision needs better search control.',5.4,7.57,9.7,.5,22,TEAL,True)
note(s,'Close the talk with the measured result and three ideas: biased selection, stochastic exploration with elite preservation, and evidence-based tuning. The frontend makes each stochastic operator visible and matches the original Python run. No claim of exact optimum or universally best settings is made.')

# 26 REFERENCES
s=new('References and reproducibility','Reference','Primary sources, local implementation, and the evidence used in this presentation')
refs=[
 ('01  Project implementation','ga_ackley.py; ga-web/engine.py; ga-web/static/','Exact operators, parameters, stage traces, and white-mode captures.',None),
 ('02  Ackley benchmark','Surjanovic & Bingham · SFU Virtual Library','Definition, landscape, recommended constants, and known optimum.','https://www.sfu.ca/~ssurjano/ackley.html'),
 ('03  Genetic algorithms','John H. Holland · Adaptation in Natural and Artificial Systems','MIT Press, 1992 edition; foundations of adaptation and genetic search.','https://doi.org/10.7551/mitpress/1090.001.0001'),
 ('04  Icons','ByteDance IconPark + the supplied custom icons','SVG assets from icons/; IconPark under Apache 2.0.','https://github.com/bytedance/IconPark'),
 ('05  Run evidence','presentation/data/run.json and summary.json','Seed-42 trace; 180 tuning runs; all 14 existing parity checks passed.',None)]
for i,(title,src,desc,url) in enumerate(refs):
    y=2.4+i*1.09
    text(s,title,.8,y,5.25,.4,22,TEAL,True)
    sh=text(s,src,6.3,y,8.7,.38,19,INK,True)
    if url:sh.text_frame.paragraphs[0].runs[0].hyperlink.address=url
    text(s,desc,6.3,y+.46,8.7,.47,17,MUTED)
    if i<4:line(s,.8,y+.96,15.05,y+.96)
text(s,'Web sources accessed 12 September 2026 · full links and run details are also in speaker notes.',.8,8.0,14.4,.3,14,MUTED)
note(s,'Full references:\n1. Local project README.md and ga_ackley.py. Web visualizer: ga-web/engine.py, server.py, static/index.html, static/js/stages.js, landscape.js, chart.js.\n2. Surjanovic, S. and Bingham, D. Virtual Library of Simulation Experiments: Test Functions and Datasets, Ackley Function. https://www.sfu.ca/~ssurjano/ackley.html\n3. Holland, John H. Adaptation in Natural and Artificial Systems: An Introductory Analysis with Applications to Biology, Control, and Artificial Intelligence. MIT Press,1992 edition. https://doi.org/10.7551/mitpress/1090.001.0001\n4. ByteDance IconPark. https://github.com/bytedance/IconPark . Apache2.0. The supplied icons/README.md identifies the custom DNA and random-initialization drawings.\n5. presentation/build_data.py generated180 runs using20 seeds42–61 for9 configurations. All measured results in presentation/data/summary.json; full walkthrough trace in run.json. verify_parity.py:14/14 exact matches. Sources accessed12 September2026.')

P.core_properties.title='2D Ackley Optimization with a Genetic Algorithm'
P.core_properties.subject='White-theme visual project presentation, seed-42 run and parameter tuning'
P.core_properties.author=''
P.core_properties.keywords='Ackley, genetic algorithm, roulette selection, crossover, mutation, elitism'
out=ROOT/'2D_Ackley_GA_Presentation.pptx';P.save(out)
(HERE/'slides.json').write_text(json.dumps(SLIDES,indent=2),encoding='utf8')
print(f'Saved {out} ({len(P.slides)} slides, {FIG} captioned figures)')
